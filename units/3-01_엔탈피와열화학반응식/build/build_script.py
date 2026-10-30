from pathlib import Path
import json,re
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
U=Path(__file__).resolve().parents[1];rows=json.loads((U/'build/lesson-content.json').read_text());D=Document();sec=D.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
 st=D.styles[name];st.font.name='Pretendard';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Pretendard');st.font.color.rgb=RGBColor(20,20,20)
normal=D.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.25;normal.paragraph_format.space_after=Pt(7)
D.styles['Title'].font.size=Pt(25);D.styles['Heading 1'].font.size=Pt(17);D.styles['Heading 2'].font.size=Pt(12);D.styles['Heading 2'].paragraph_format.space_before=Pt(12);D.styles['Heading 2'].paragraph_format.space_after=Pt(5)
def paragraph(text,style=None):
 p=D.add_paragraph(style=style)
 # Chemical-state parentheses stay upright; only Latin state letters italic.
 start=0
 for m in re.finditer(r'\((s|l|g|aq)\)',text):
  p.add_run(text[start:m.start()+1]);p.add_run(m.group(1)).italic=True;p.add_run(')');start=m.end()
 p.add_run(text[start:]);return p
paragraph('엔탈피와 열화학 반응식 수업 대본','Title')
paragraph('물질과 에너지 3단원 첫 번째 소단원')
paragraph('전체 31장 · 수업 50분 · 천재교육 2022 개정 교과서 74~81쪽')
paragraph('수업 운영','Heading 1')
paragraph('손난로 관찰에서 시작하여 열화학 반응식, 연료 비교, 결합 에너지 계산으로 이어지는 수업이다. 각 장의 실제 발화를 읽고, 그림을 가리키는 순서와 오개념 점검을 함께 활용한다. 제시된 시간에는 학생의 예상, 계산, 짝 설명 시간을 포함한다.')
paragraph('진행 순서','Heading 2')
for t in ['1~4장 도입과 손난로 예상 4분','5~13장 엔탈피와 열화학 반응식 15분','14~21장 연료의 효율성 13분','22~29장 결합 에너지와 반응 엔탈피 14분','30~31장 정리와 형성평가 4분']:paragraph(t)
paragraph('확인 문제와 답 공개','Heading 2')
paragraph('13장은 엔탈피 비교 답과 암모니아 반응식을 차례로, 21장은 계산값과 비교 근거를 차례로, 31장은 HF 계산 답을 공개한다. PPT 슬라이드 쇼에서는 클릭으로 진행한다. 전체 미리보기의 답 보기 버튼으로도 확인할 수 있다.')
paragraph('자료의 해석','Heading 2')
paragraph('교과서 사진과 입자 모형은 출처를 각 장에 기록했다. 엔탈피 수준 그래프는 교과서 개념을 편집 가능한 도형으로 재구성했다. 생성형 학생 삽화는 설명을 돕는 장면이며 실험 관찰 결과가 아니다. 평균 결합 에너지 계산값은 기체 반응의 근삿값이고, 방출열의 크기와 반응 엔탈피의 음수를 구별한다.')
for r in rows:
 D.add_page_break();paragraph('Slide '+str(r['number']),'Heading 1');paragraph(r['title'],'Heading 1');paragraph(f"권장 시간 {r['minutes']:g}분")
 paragraph('먼저 알아둘 점','Heading 2');paragraph(r['background'])
 paragraph('발표 대본','Heading 2');paragraph(r['script'])
 paragraph('화면에서 가리킬 순서','Heading 2');paragraph(r['point'])
 paragraph('오개념 점검','Heading 2');paragraph(r['caution'])
 paragraph('다음 장으로 연결','Heading 2');paragraph(r['next'])
 p=paragraph('출처: '+r['source']);p.paragraph_format.space_before=Pt(12)
 for run in p.runs:run.font.size=Pt(9);run.font.color.rgb=RGBColor(90,90,90)
for el in list(D.styles.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
for el in list(D.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
for el in D.styles.element.iter(qn('w:rFonts')):
 for a in list(el.attrib):
  if 'Theme' in a: del el.attrib[a]
D.core_properties.title='엔탈피와 열화학 반응식 수업 대본';D.core_properties.author='사이음';D.save(U/'build/full-v1/script.docx');print('DOCX saved with31 slide sections')
