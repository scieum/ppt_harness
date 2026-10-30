from pathlib import Path
import json,re
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
U=Path(__file__).resolve().parents[1];rows=json.loads((U/'build/full-v2/slides.json').read_text());D=Document();sec=D.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
 st=D.styles[name];st.font.name='Pretendard';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Pretendard');st.font.color.rgb=RGBColor(20,20,20)
normal=D.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.25;normal.paragraph_format.space_after=Pt(7)
D.styles['Title'].font.size=Pt(25);D.styles['Heading 1'].font.size=Pt(17);D.styles['Heading 2'].font.size=Pt(12);D.styles['Heading 2'].paragraph_format.space_before=Pt(12);D.styles['Heading 2'].paragraph_format.space_after=Pt(5)
def paragraph(text,style=None):
 p=D.add_paragraph(style=style)
 # State letters and enthalpy H are italic; chemical H and parentheses upright.
 italic=set()
 for m in re.finditer(r'\((s|l|g|aq)\)',text):italic.update(range(m.start()+1,m.end()-1))
 for m in re.finditer('ΔH',text):italic.add(m.start()+1)
 if globals().get('r',{}).get('number') in [8,9]:
  italic.update(m.start() for m in re.finditer('H',text))
 start=0
 while start<len(text):
  flag=start in italic;end=start+1
  while end<len(text) and (end in italic)==flag:end+=1
  p.add_run(text[start:end]).italic=flag;start=end
 return p
paragraph('엔탈피와 열화학 반응식 수업 대본','Title')
paragraph('물질과 에너지 3단원 첫 번째 소단원')
paragraph('전체 39장 · 수업 63분 · 천재교육 2022 개정 교과서 74~81쪽')
paragraph('수업 운영','Heading 1')
paragraph('손난로 관찰에서 시작하여 발열·흡열 반응과 엔탈피, 열화학 반응식, 연료 비교, 결합 에너지 계산으로 이어지는 수업이다. 각 장의 실제 발화를 읽고, 그림을 가리키는 순서와 오개념 점검을 함께 활용한다. 제시된 시간에는 학생의 예상, 계산, 짝 설명 시간을 포함한다.')
paragraph('진행 순서','Heading 2')
for t in ['1~4장 도입과 손난로 예상 4분','5~11장 엔탈피 11.5분','12~21장 열화학 반응식 16.5분','22~29장 연료의 효율성 13분','30~37장 결합 에너지와 반응 엔탈피 14분','38~39장 정리와 형성평가 4분']:paragraph(t)
paragraph('확인 문제와 답 공개','Heading 2')
paragraph('21장은 엔탈피 비교 답과 암모니아 반응식을 차례로, 29장은 계산값과 비교 근거를 차례로, 39장은 HF 계산 답을 공개한다. PPT 슬라이드 쇼에서는 클릭으로 진행한다. 전체 미리보기의 답 보기 버튼으로도 확인할 수 있다.')
paragraph('자료의 해석','Heading 2')
paragraph('교과서 사진과 입자 모형은 출처를 각 장에 기록했다. 엔탈피 그래프, 반응식과 결합 모형은 교과서·지도서 원본을 우선 사용했다. 연료 비교 표는 원본 수치를 유지해 편집 가능한 표로 옮기고, 양수로 제시한 값의 항목명을 방출열로 명확히 했다. 보충 반응식과 요약표는 직접 작성했다. 생성형 학생 삽화는 설명을 돕는 장면이며 실험 관찰 결과가 아니다. 평균 결합 에너지 계산값은 기체 반응의 근삿값이고, 방출열의 크기와 반응 엔탈피의 음수를 구별한다.')
for r in rows:
 D.add_page_break();paragraph('Slide '+str(r['number']),'Heading 1');paragraph(r['title'],'Heading 1');paragraph(f"권장 시간 {r['minutes']:g}분")
 if r['background'] != r['caution']:
  paragraph('먼저 알아둘 점','Heading 2');paragraph(r['background'])
 paragraph('발표 대본','Heading 2');paragraph(r['script'])
 paragraph('화면에서 가리킬 순서','Heading 2');paragraph(r['point'])
 paragraph('오개념 점검','Heading 2');paragraph(r['caution'])
 paragraph('다음 장으로 연결','Heading 2');paragraph(r['next'])
 p=paragraph('출처: '+r['source']);p.paragraph_format.space_before=Pt(12)
 for run in p.runs:run.font.size=Pt(9);run.font.color.rgb=RGBColor(90,90,90)
for el in list(D.styles.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
for el in list(D.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
for el in D.styles.element.iter(qn('w:rFonts')):
 for a in list(el.attrib):
  if 'Theme' in a: del el.attrib[a]
D.core_properties.title='엔탈피와 열화학 반응식 수업 대본';D.core_properties.author='사이음';D.save(U/'build/full-v2/script.docx');print('DOCX saved with39 slide sections')
